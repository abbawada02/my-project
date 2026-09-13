from decimal import Decimal
from io import BytesIO

from django.test import TestCase
from openpyxl import load_workbook
from rest_framework.test import APIRequestFactory, force_authenticate

from accounts.models import AdminProfile, CustomUser, LecturerProfile, StudentProfile
from courses.models import AcademicSession, Course, CourseAllocation, Enrollment, Programme
from .models import Result, SemesterGPA, SemesterResultBatch
from .signals import recompute_gpa
from .views import (
	ApproveBatchView,
	LecturerResultExcelUploadView,
	LecturerResultUploadView,
	ScoreSheetTemplateView,
	StudentGPAView,
	StudentResultsView,
)


class ResultWorkflowTests(TestCase):
	def setUp(self):
		self.factory = APIRequestFactory()
		self.session = AcademicSession.objects.create(
			name="2099/2100", start_date="2099-09-01", end_date="2100-08-31", is_current=True
		)
		self.programme = Programme.objects.create(name="Computer Science", programme_type="MSC")
		self.admin = CustomUser.objects.create_user(
			email="admin@test.com", password="password123", role="ADMIN", first_name="PG", last_name="Coordinator"
		)
		AdminProfile.objects.create(user=self.admin, admin_role="PG_COORDINATOR")
		self.lecturers = []
		for index in range(3):
			user = CustomUser.objects.create_user(
				email=f"lecturer{index}@test.com", password="password123", role="LECTURER",
				first_name=f"Lecturer{index}", last_name="Test",
			)
			self.lecturers.append(LecturerProfile.objects.create(
				user=user, staff_id=f"STAFF-{index}", status="ACTIVE"
			))
		student_user = CustomUser.objects.create_user(
			email="student@test.com", password="password123", role="STUDENT",
			first_name="Student", last_name="Test",
		)
		self.student = StudentProfile.objects.create(
			user=student_user, programme="MSC", matric_number="MSC-001", admission_year=2099, status="ACTIVE"
		)
		self.allocations = []
		for index, lecturer in enumerate(self.lecturers, start=1):
			course = Course.objects.create(
				code=f"CSC9{index:02d}", title=f"Test Course {index}", credit_units=3,
				semester="FIRST", programme=self.programme,
			)
			allocation = CourseAllocation.objects.create(
				lecturer=lecturer, course=course, session=self.session, semester="FIRST"
			)
			self.allocations.append(allocation)
			Enrollment.objects.create(student=self.student, allocation=allocation, status="APPROVED")
		self.batch = SemesterResultBatch.objects.create(
			session=self.session, semester="FIRST", programme_type="MSC"
		)

	def upload_json(self, lecturer, enrollment, score="70"):
		request = self.factory.post(
			"/api/results/upload/",
			{"batch": self.batch.id, "results": [{"enrollment": enrollment.id, "score": score}]},
			format="json",
		)
		force_authenticate(request, user=lecturer.user)
		return LecturerResultUploadView.as_view()(request)

	def downloaded_template(self, allocation):
		request = self.factory.get(f"/api/results/template/{allocation.id}/")
		force_authenticate(request, user=allocation.lecturer.user)
		response = ScoreSheetTemplateView.as_view()(request, allocation_pk=allocation.id)
		self.assertEqual(response.status_code, 200)
		return load_workbook(BytesIO(response.content))

	def upload_workbook(self, lecturer, allocation, workbook):
		file = BytesIO()
		workbook.save(file)
		file.seek(0)
		file.name = "CSC_TEST_SCORE_SHEET.xlsx"
		request = self.factory.post(
			f"/api/results/upload-excel/{allocation.id}/",
			{"file": file},
			format="multipart",
		)
		force_authenticate(request, user=lecturer.user)
		return LecturerResultExcelUploadView.as_view()(request, allocation_pk=allocation.id)

	def approve(self):
		request = self.factory.post(
			f"/api/results/batches/{self.batch.id}/approve/", {"action": "approve"}, format="json"
		)
		force_authenticate(request, user=self.admin)
		return ApproveBatchView.as_view()(request, pk=self.batch.id)

	def test_multiple_lecturers_share_one_pending_batch(self):
		for lecturer, allocation in zip(self.lecturers, self.allocations):
			response = self.upload_json(lecturer, allocation.enrollments.get())
			self.assertEqual(response.status_code, 200)
		self.assertEqual(Result.objects.filter(batch=self.batch).count(), 3)
		self.assertEqual(SemesterResultBatch.objects.count(), 1)

	def test_incomplete_batch_cannot_be_approved(self):
		self.upload_json(self.lecturers[0], self.allocations[0].enrollments.get())
		response = self.approve()
		self.assertEqual(response.status_code, 400)
		self.batch.refresh_from_db()
		self.assertEqual(self.batch.status, SemesterResultBatch.Status.PENDING)
		self.assertEqual(response.data["progress"]["missing_course_count"], 2)

	def test_all_required_courses_submitted_allows_approval(self):
		for lecturer, allocation in zip(self.lecturers, self.allocations):
			self.upload_json(lecturer, allocation.enrollments.get())
		response = self.approve()
		self.assertEqual(response.status_code, 200)
		self.batch.refresh_from_db()
		self.assertEqual(self.batch.status, SemesterResultBatch.Status.APPROVED)

	def test_approved_batch_rejects_further_uploads(self):
		for lecturer, allocation in zip(self.lecturers, self.allocations):
			self.upload_json(lecturer, allocation.enrollments.get())
		self.assertEqual(self.approve().status_code, 200)
		response = self.upload_json(self.lecturers[0], self.allocations[0].enrollments.get(), score="80")
		self.assertEqual(response.status_code, 400)
		self.assertIn("approved", response.data["detail"].lower())

	def test_only_approved_results_contribute_to_gpa(self):
		enrollment = self.allocations[0].enrollments.get()
		self.upload_json(self.lecturers[0], enrollment)
		self.assertFalse(SemesterGPA.objects.filter(student=self.student).exists())
		self.batch.status = SemesterResultBatch.Status.APPROVED
		self.batch.save()
		recompute_gpa(self.student, self.session, "FIRST")
		gpa = SemesterGPA.objects.get(student=self.student, session=self.session, semester="FIRST")
		self.assertEqual(gpa.gpa, Decimal("5.00"))

	def test_student_can_retrieve_approved_results(self):
		enrollment = self.allocations[0].enrollments.get()
		self.upload_json(self.lecturers[0], enrollment)
		self.batch.status = SemesterResultBatch.Status.APPROVED
		self.batch.save()
		request = self.factory.get("/api/results/mine/")
		force_authenticate(request, user=self.student.user)
		response = StudentResultsView.as_view()(request)
		self.assertEqual(response.status_code, 200)
		self.assertEqual(len(response.data), 1)

	def test_lecturer_cannot_upload_another_lecturers_allocation(self):
		enrollment = self.allocations[1].enrollments.get()
		response = self.upload_json(self.lecturers[0], enrollment)
		self.assertEqual(response.status_code, 403)
		self.assertFalse(Result.objects.filter(enrollment=enrollment).exists())

	def test_student_gpa_endpoint_returns_gpa_and_cgpa_after_approval(self):
		enrollment = self.allocations[0].enrollments.get()
		self.upload_json(self.lecturers[0], enrollment)
		result = Result.objects.get(enrollment=enrollment)
		self.batch.status = SemesterResultBatch.Status.APPROVED
		self.batch.save()
		result.save()
		request = self.factory.get("/api/results/gpa/")
		force_authenticate(request, user=self.student.user)
		response = StudentGPAView.as_view()(request)
		self.assertEqual(response.status_code, 200)
		self.assertEqual(len(response.data["semester_gpas"]), 1)
		self.assertIsNotNone(response.data["cumulative"])

	def test_downloaded_template_with_scores_uploads_successfully(self):
		allocation = self.allocations[0]
		workbook = self.downloaded_template(allocation)
		worksheet = workbook.active
		worksheet.cell(row=10, column=4, value=25)
		worksheet.cell(row=10, column=5, value=50)
		response = self.upload_workbook(allocation.lecturer, allocation, workbook)
		self.assertEqual(response.status_code, 200)
		self.assertEqual(response.data["uploaded"], 1)
		result = Result.objects.get(enrollment=allocation.enrollments.get())
		self.assertEqual(result.score, Decimal("75.00"))

	def test_invalid_score_is_reported(self):
		allocation = self.allocations[0]
		workbook = self.downloaded_template(allocation)
		worksheet = workbook.active
		worksheet.cell(row=10, column=4, value=80)
		worksheet.cell(row=10, column=5, value=30)
		response = self.upload_workbook(allocation.lecturer, allocation, workbook)
		self.assertEqual(response.status_code, 200)
		self.assertEqual(response.data["uploaded"], 0)
		self.assertIn("exceeds 100", response.data["errors"][0]["detail"])

	def test_wrong_course_metadata_is_rejected(self):
		allocation = self.allocations[0]
		workbook = self.downloaded_template(allocation)
		workbook.active.cell(row=5, column=1, value="CSC999 : Wrong Course")
		workbook.active.cell(row=8, column=1, value="CSC999")
		response = self.upload_workbook(allocation.lecturer, allocation, workbook)
		self.assertEqual(response.status_code, 400)
		self.assertIn("Wrong course", response.data["detail"])

	def test_approved_batch_rejects_excel_upload(self):
		allocation = self.allocations[0]
		workbook = self.downloaded_template(allocation)
		workbook.active.cell(row=10, column=4, value=25)
		workbook.active.cell(row=10, column=5, value=50)
		self.batch.status = SemesterResultBatch.Status.APPROVED
		self.batch.save()
		response = self.upload_workbook(allocation.lecturer, allocation, workbook)
		self.assertEqual(response.status_code, 400)
		self.assertIn("approved", response.data["detail"].lower())
